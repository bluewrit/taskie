"""End-to-end API tests: tasks, file upload/preview, agent chat/plan/eval."""
import json
import os
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

import pytest

# isolate the database before importing the app
_TMP = tempfile.mkdtemp()
os.environ["TASKIE_DATA_DIR"] = _TMP
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

client = TestClient(app)


def _create_task(title="Test task", **kw):
    payload = {"title": title, "priority": kw.pop("priority", "medium"),
               "due_date": kw.pop("due_date", None), **kw}
    r = client.post("/api/tasks", json=payload)
    assert r.status_code == 201, r.text
    return r.json()


# ------------------------------------------------------------------ tasks
def test_health():
    assert client.get("/api/health").json()["status"] == "ok"


def test_task_crud():
    t = _create_task("Write tests", priority="high")
    tid = t["id"]
    assert client.get(f"/api/tasks/{tid}").json()["title"] == "Write tests"

    r = client.put(f"/api/tasks/{tid}", json={"status": "in_progress", "progress": 50})
    assert r.status_code == 200
    assert r.json()["status"] == "in_progress"

    r = client.post(f"/api/tasks/{tid}/complete")
    assert r.json()["status"] == "done"
    assert r.json()["progress"] == 100
    assert r.json()["completed_at"] is not None

    assert client.delete(f"/api/tasks/{tid}").status_code == 204
    assert client.get(f"/api/tasks/{tid}").status_code == 404


def test_task_validation():
    assert client.post("/api/tasks", json={"title": ""}).status_code == 422
    assert client.post("/api/tasks", json={"title": "x", "status": "bogus"}).status_code == 422


# ------------------------------------------------------------------ files
def test_upload_and_text_preview():
    t = _create_task("Files task")
    r = client.post(f"/api/tasks/{t['id']}/files",
                    files={"file": ("hello.txt", b"hello world", "text/plain")})
    assert r.status_code == 201, r.text
    fid = r.json()["id"]
    assert r.json()["size"] == 11

    prev = client.get(f"/api/files/{fid}/preview").json()
    assert prev["kind"] == "text"
    assert prev["content"] == "hello world"

    raw = client.get(f"/api/files/{fid}/raw")
    assert raw.status_code == 200
    assert raw.content == b"hello world"
    assert raw.headers["accept-ranges"] == "bytes"


def test_upload_csv_preview_is_table():
    t = _create_task("CSV task")
    csv_bytes = b"name,qty\napple,3\npear,7\n"
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("stock.csv", csv_bytes, "text/csv")}).json()["id"]
    prev = client.get(f"/api/files/{fid}/preview").json()
    assert prev["kind"] == "table"
    assert prev["columns"] == ["name", "qty"]
    assert prev["rows"][1] == ["pear", "7"]


def test_upload_docx_preview_extracts_text():
    import io
    import docx
    doc = docx.Document()
    doc.add_paragraph("Alpha content marker")
    buf = io.BytesIO()
    doc.save(buf)

    t = _create_task("Docx task")
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("doc.docx", buf.getvalue(),
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
                      ).json()["id"]
    prev = client.get(f"/api/files/{fid}/preview").json()
    assert prev["kind"] == "document"
    assert "Alpha content marker" in prev["content"]


def test_upload_xlsx_preview_is_sheets():
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws.append(["a", "b"])
    ws.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)

    t = _create_task("Xlsx task")
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("book.xlsx", buf.getvalue(),
                                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                      ).json()["id"]
    prev = client.get(f"/api/files/{fid}/preview").json()
    assert prev["kind"] == "sheets"
    assert prev["sheets"][0]["name"] == "S1"
    assert prev["sheets"][0]["columns"] == ["a", "b"]


def test_upload_zip_preview_lists_entries():
    import io
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("inner/a.txt", "hi")
    t = _create_task("Zip task")
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("bundle.zip", buf.getvalue(), "application/zip")}).json()["id"]
    prev = client.get(f"/api/files/{fid}/preview").json()
    assert prev["kind"] == "archive"
    assert any(e["name"].startswith("inner/") for e in prev["entries"])


def test_range_request_returns_206():
    t = _create_task("Range task")
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("blob.bin", bytes(range(256)), "application/octet-stream")}).json()["id"]
    r = client.get(f"/api/files/{fid}/raw", headers={"Range": "bytes=10-19"})
    assert r.status_code == 206
    assert r.content == bytes(range(10, 20))
    assert r.headers["content-range"] == "bytes 10-19/256"


def test_download_forces_attachment():
    t = _create_task("Download task")
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("data.bin", b"\x00\x01", "application/octet-stream")}).json()["id"]
    r = client.get(f"/api/files/{fid}/download")
    assert "attachment" in r.headers["content-disposition"]


def test_delete_file():
    t = _create_task("Delete file task")
    fid = client.post(f"/api/tasks/{t['id']}/files",
                      files={"file": ("x.txt", b"x", "text/plain")}).json()["id"]
    assert client.delete(f"/api/files/{fid}").status_code == 204
    assert client.get(f"/api/files/{fid}/preview").status_code == 404


# ------------------------------------------------------------------ agent
def test_agent_create_task_via_chat():
    r = client.post("/api/agent/chat", json={
        "message": 'add task "Prepare investor deck" due tomorrow high priority'})
    assert r.status_code == 200, r.text
    body = r.json()
    created = [a for a in body["actions"] if a["tool"] == "create_task"]
    assert created and created[0]["result"]["created"]
    assert created[0]["result"]["priority"] == "high"
    assert created[0]["result"]["due_date"] == (date.today() + timedelta(days=1)).isoformat()


def test_agent_create_task_bare_weekday_due_date():
    from datetime import date as _date, timedelta as _td
    r = client.post("/api/agent/chat", json={
        "message": 'add task "Prepare keynote slides" due friday critical priority'})
    created = [a for a in r.json()["actions"] if a["tool"] == "create_task"][0]["result"]
    assert created["created"] and created["priority"] == "critical"
    # resolve the expected next Friday the same way the NLU does
    today = _date.today()
    expected = today + _td(days=(4 - today.weekday()) % 7 or 7)
    assert created["due_date"] == expected.isoformat()
    client.delete(f"/api/tasks/{created['task_id']}")  # cleanup


def test_agent_complete_task_via_chat():
    t = _create_task("Finish the quarterly report")
    r = client.post("/api/agent/chat", json={"message": f"mark task #{t['id']} done"})
    assert any(a["tool"] == "complete_task" and a["result"].get("completed") for a in r.json()["actions"])
    assert client.get(f"/api/tasks/{t['id']}").json()["status"] == "done"


def test_agent_recommendations_and_apply():
    _create_task("Overdue urgent thing", priority="critical",
                 due_date=(date.today() - timedelta(days=3)).isoformat())
    r = client.get("/api/agent/recommendations")
    recs = r.json()["recommendations"]
    assert any(rec["type"] == "overdue" for rec in recs)
    # scores were persisted
    tasks = client.get("/api/tasks", params={"status": "todo"}).json()
    assert any(t["agent_score"] > 0 for t in tasks)

    r2 = client.get("/api/agent/recommendations", params={"apply": True})
    assert r2.status_code == 200
    assert isinstance(r2.json()["applied"], list)


def test_agent_plan():
    r = client.get("/api/agent/plan", params={"horizon": 5})
    plan = r.json()
    assert plan["horizon_days"] == 5
    assert len(plan["schedule"]) == 5
    assert plan["summary"]["tasks_planned"] >= 1


def test_agent_evaluation():
    r = client.get("/api/agent/evaluation")
    ev = r.json()
    assert ev["grade"] in list("ABCDF")
    assert 0 <= ev["overall_score"] <= 100
    assert "metrics" in ev and "throughput" in ev
    assert len(ev["throughput"]) == 14


def test_agent_actions_logged():
    r = client.get("/api/agent/actions")
    kinds = {a["kind"] for a in r.json()}
    assert "chat" in kinds
